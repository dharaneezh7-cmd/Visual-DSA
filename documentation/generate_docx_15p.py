import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

def add_page_border(section):
    secPr = section._sectPr
    pgBorders = parse_xml(
        f'<w:pgBorders {nsdecls("w")} w:offsetFrom="page">'
        f'  <w:top w:val="single" w:sz="12" w:space="24" w:color="000000"/>'
        f'  <w:left w:val="single" w:sz="12" w:space="24" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="12" w:space="24" w:color="000000"/>'
        f'  <w:right w:val="single" w:sz="12" w:space="24" w:color="000000"/>'
        f'</w:pgBorders>'
    )
    secPr.append(pgBorders)

def create_15_page_document():
    doc = Document()
    
    # Page setup - A4 with 0.75 inch margins
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    add_page_border(section)

    # Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(3)

    def add_heading(text, space_before=6, space_after=4):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(12)
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        return p

    # =========================================================================
    # PAGE 1: COVER PAGE
    # =========================================================================
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(25)
    p_logo.paragraph_format.space_after = Pt(15)
    
    cit_logo_path = 'documentation/extracted_p1_img0_Image18.png'
    if os.path.exists(cit_logo_path):
        run_logo = p_logo.add_run()
        run_logo.add_picture(cit_logo_path, width=Inches(1.5))

    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("COIMBATORE INSTITUTE OF TECHNOLOGY—641014\nMSC – SOFTWARE SYSTEMS.")
    r_inst.bold = True
    r_inst.font.size = Pt(14)
    p_inst.paragraph_format.space_after = Pt(25)

    p_proj = doc.add_paragraph()
    p_proj.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_proj = p_proj.add_run("PROJECT TITLE : VISUAL DSA — INTERACTIVE DATA STRUCTURES AND ALGORITHMS VISUALIZATION PLATFORM.")
    r_proj.bold = True
    r_proj.font.size = Pt(13)
    p_proj.paragraph_format.space_after = Pt(8)

    p_mod = doc.add_paragraph()
    p_mod.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_mod = p_mod.add_run("MODULE : Interactive Memory Visualizer, Pseudocode Synchronization & AI Learning Assistant.")
    r_mod.font.size = Pt(12)
    p_mod.paragraph_format.space_after = Pt(45)

    # Student Details Table
    table_p1 = doc.add_table(rows=6, cols=2)
    table_p1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_p1, color="000000", sz="6")

    meta_rows = [
        ("NAME", "DHARANEESH V"),
        ("REGISTER NUMBER", "2503717672621011"),
        ("COURSE", "MSC – SOFTWARE SYSTEMS"),
        ("SUBJECT CODE", "25SS322"),
        ("SUBJECT", "PYTHON PROGRAMMING LABORATORY"),
        ("SEMESTER", "3")
    ]

    col_widths = [Inches(2.5), Inches(3.8)]
    for i, (k, v) in enumerate(meta_rows):
        row = table_p1.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = col_widths[0]
        c1.width = col_widths[1]
        set_cell_margins(c0, top=100, bottom=100, left=150, right=150)
        set_cell_margins(c1, top=100, bottom=100, left=150, right=150)
        
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(k)
        r0.bold = True
        r0.font.size = Pt(10.5)

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(v)
        r1.bold = True
        r1.font.size = Pt(10.5)

    doc.add_page_break()

    # =========================================================================
    # PAGE 2: SYNOPSIS, OBJECTIVE, AND SCOPE
    # =========================================================================
    add_heading("SYNOPSIS", space_before=4, space_after=3)
    p_syn1 = doc.add_paragraph(
        "Visual DSA is an interactive educational web-based platform developed to simplify the process of learning "
        "fundamental Data Structures and Algorithms. The platform replaces passive, static textbook code with responsive, "
        "step-by-step visual memory simulations, allowing learners to inspect pointer mutations, element displacements, "
        "and recursion mechanics in real time."
    )
    p_syn1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p_syn2 = doc.add_paragraph(
        "Module: Interactive Memory Visualizer, Pseudocode Synchronization & AI Learning Assistant is responsible for "
        "the core experiential learning lifecycle. It delivers dedicated visual sandboxes for linear arrays, singly linked "
        "lists, stacks (LIFO), queues (FIFO), circular queues (ring buffers), linear/binary searching, and five distinct "
        "sorting algorithms. Synchronized line-by-line pseudocode execution dynamically coordinates with graphic canvas state changes."
    )
    p_syn2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p_syn3 = doc.add_paragraph(
        "The application uses React 19 with Vite and TypeScript for the frontend client interface, and Node.js with "
        "Express.js for the primary API gateway. MongoDB is utilized for persistent cloud storage of student learning "
        "progress, operational activities, and quiz practice metrics. A supporting Python FastAPI microservice computes "
        "statistical mastery analytics and delivers SMTP password-reset OTP emails. An integrated Ollama-powered mascot tutor "
        "('Lemmy') provides context-aware tutoring and audio/speech synthesis."
    )
    p_syn3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    add_heading("OBJECTIVE", space_before=6, space_after=3)
    objectives = [
        "Render interactive, step-by-step memory visualizers for linear data structures, searching, and sorting.",
        "Synchronize algorithmic state transformations with highlighted, line-by-line pseudocode execution.",
        "Demonstrate pointer mutations, head/tail allocations, and modulo ring-buffer wrap-arounds graphically.",
        "Provide formative retrieval practice quizzes with immediate explanation reveals and accuracy metrics.",
        "Enforce defensive REST API security using JWT session tokens, bcrypt password hashing, and rate limiters.",
        "Deliver automated analytical diagnostic reports identifying student strengths, weaknesses, and study recommendations.",
        "Integrate a conversational local LLM AI tutor ('Lemmy') providing Big-O breakdowns and real-world analogies.",
        "Generate verifiable, high-resolution HTML5 Canvas Certificates of Completion with PNG and PDF exports."
    ]
    for obj in objectives:
        p_obj = doc.add_paragraph(obj, style='List Bullet')
        p_obj.paragraph_format.space_after = Pt(1.5)
        p_obj.paragraph_format.line_spacing = 1.1

    add_heading("SCOPE", space_before=6, space_after=3)
    scopes = [
        "Interactive visual sandboxes covering Arrays, Singly Linked Lists, Stacks, Queues, Circular Queues, Searching, and Sorting.",
        "Synchronous visualizer speed-control loop supporting custom operation parameters and boundary condition alerts.",
        "Dual-layer state persistence supporting authenticated MongoDB storage and seamless unauthenticated guest localStorage.",
        "FastAPI microservice executing statistical aggregation, mastery level classification, and SMTP email delivery.",
        "REST API gateway exposing validated, rate-limited endpoints for authentication, progress tracking, and quiz submissions.",
        "Client-side HTML5 Canvas rendering engine producing customized, verifiable completion certificates with unique serial IDs."
    ]
    for sc in scopes:
        p_sc = doc.add_paragraph(sc, style='List Bullet')
        p_sc.paragraph_format.space_after = Pt(1.5)
        p_sc.paragraph_format.line_spacing = 1.1

    doc.add_page_break()

    # =========================================================================
    # PAGE 3: PROJECT DESCRIPTION AND DB DESIGN
    # =========================================================================
    add_heading("PROJECT DESCRIPTION", space_before=2, space_after=2)
    p_pd1 = doc.add_paragraph(
        "Visual DSA is a full-stack, distributed web application consisting of a React 19 (Vite) frontend, a Node.js/Express.js "
        "API gateway, a MongoDB database, and a supporting Python FastAPI microservice. The core visualizer module delivers "
        "the complete interactive algorithmic learning lifecycle, connecting interactive simulations with verified cloud mastery records."
    )
    p_pd1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_pd1.paragraph_format.space_after = Pt(2)

    p_pd2 = doc.add_paragraph(
        "The frontend provides an intuitive multi-tab visualizer with playback speed sliders, synchronized pseudocode blocks, "
        "and conversational AI assistance. The Node.js backend manages secure REST API routing, input sanitization, rate limiting, "
        "and Mongoose persistence, while communicating with the Python microservice for heavy statistical analytics."
    )
    p_pd2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_pd2.paragraph_format.space_after = Pt(4)

    # Technology Table (6 rows)
    table_tech = doc.add_table(rows=7, cols=2)
    table_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_tech, color="D3D3D3", sz="4")

    h0, h1 = table_tech.rows[0].cells[0], table_tech.rows[0].cells[1]
    h0.paragraphs[0].add_run("Technology").bold = True
    h1.paragraphs[0].add_run("Purpose").bold = True
    set_cell_background(h0, "F1F5F9")
    set_cell_background(h1, "F1F5F9")
    set_cell_margins(h0, top=50, bottom=50, left=100, right=100)
    set_cell_margins(h1, top=50, bottom=50, left=100, right=100)

    tech_data = [
        ("React 19 & TypeScript", "Frontend user interface, component hierarchy, typed state models, and reactive DOM rendering"),
        ("Vite (v8)", "Fast Hot Module Replacement (HMR) build tooling and optimized production bundling"),
        ("Node.js & Express (v4)", "Asynchronous runtime executing REST API gateway, rate limiting, and security middleware"),
        ("MongoDB & Mongoose", "NoSQL document database and Object Data Modeling (ODM) defining schemas and validation"),
        ("Python 3 & FastAPI", "Supporting microservice for statistical mastery analytics, study time formatting, and SMTP OTP"),
        ("Ollama / Local LLM", "Contextual algorithmic tutor ('Lemmy') powered by llama3.2:3b with Web Audio and Speech synthesis")
    ]

    for idx, (t, p) in enumerate(tech_data, start=1):
        row = table_tech.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.5)
        run_t = c0.paragraphs[0].add_run(t)
        run_t.bold = True
        run_t.font.size = Pt(9.5)
        run_p = c1.paragraphs[0].add_run(p)
        run_p.font.size = Pt(9.5)
        set_cell_margins(c0, top=40, bottom=40, left=90, right=90)
        set_cell_margins(c1, top=40, bottom=40, left=90, right=90)

    add_heading("DB DESIGN", space_before=5, space_after=2)
    p_dbd = doc.add_paragraph(
        "Database Design describes how data is organized and stored in MongoDB. The system primarily interacts with the "
        "'users', 'progresses', 'practices', and 'activities' collections. The primary Progress schema structure is detailed below:"
    )
    p_dbd.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_dbd.paragraph_format.space_after = Pt(3)

    # DB Schema Table (7 data rows)
    table_db = doc.add_table(rows=8, cols=3)
    table_db.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_db, color="D3D3D3", sz="4")

    db_headers = ("Field", "Data Type", "Description")
    for j, dh in enumerate(db_headers):
        cell = table_db.rows[0].cells[j]
        cell.paragraphs[0].add_run(dh).bold = True
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=50, bottom=50, left=80, right=80)

    schema_data = [
        ("_id", "ObjectId", "Unique progress document identifier generated by MongoDB"),
        ("userId", "ObjectId", "Reference foreign key linking to the authenticated User document"),
        ("topic", "String", "Curriculum domain category (e.g., 'Array', 'Stack', 'CircularQueue')"),
        ("concept", "String", "Specific module concept name (e.g., 'Operations', 'Traversal')"),
        ("completed", "Boolean", "Completion flag indicating full mastery of the topic operations"),
        ("completionPercentage", "Number", "Calculated mastery progress metric ranging from 0 to 100"),
        ("timeSpent", "Number", "Cumulative study and visualizer engagement time in seconds")
    ]

    for idx, (f, dt, desc) in enumerate(schema_data, start=1):
        row = table_db.rows[idx]
        c0, c1, c2 = row.cells[0], row.cells[1], row.cells[2]
        c0.width = Inches(1.8)
        c1.width = Inches(1.2)
        c2.width = Inches(3.7)
        r0 = c0.paragraphs[0].add_run(f)
        r0.bold = True
        r0.font.size = Pt(9)
        r1 = c1.paragraphs[0].add_run(dt)
        r1.font.size = Pt(9)
        r2 = c2.paragraphs[0].add_run(desc)
        r2.font.size = Pt(9)
        set_cell_margins(c0, top=35, bottom=35, left=70, right=70)
        set_cell_margins(c1, top=35, bottom=35, left=70, right=70)
        set_cell_margins(c2, top=35, bottom=35, left=70, right=70)

    p_dbflow = doc.add_paragraph()
    p_dbflow.paragraph_format.space_before = Pt(4)
    p_dbflow.paragraph_format.space_after = Pt(2)
    r_flow = p_dbflow.add_run("Database flow: React Frontend -> Progress API (POST /api/progress) -> Node.js/Express -> Mongoose -> MongoDB -> Progresses & Activities Collections.")
    r_flow.bold = True
    r_flow.font.size = Pt(8.5)

    p_atom = doc.add_paragraph(
        "Data persistence guarantees consistency: upon visualizer operation completion, student records are atomically "
        "upserted using findOneAndUpdate() with compound index { userId: 1, topic: 1, concept: 1 }, cumulative time tallies are "
        "incremented, and verified activity audit events are logged in the Activity collection."
    )
    p_atom.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_atom.paragraph_format.line_spacing = 1.1
    p_atom.paragraph_format.space_after = Pt(2)

    doc.add_page_break()

    # =========================================================================
    # PAGES 4 to 8: DIAGRAMS (1 Page Each)
    # =========================================================================
    diagram_pages = [
        ("DATA FLOW DIAGRAM (DFD)",
         "A Data Flow Diagram (DFD) represents how data moves through the Visual DSA interactive platform. It illustrates "
         "the interaction between learners, frontend visualizers, validation layers, Express API endpoints, MongoDB collections, "
         "the Python analysis service, and the local AI assistant runtime.",
         "documentation/diagrams/page4_dfd.png"),

        ("ER DIAGRAM",
         "An Entity Relationship (ER) Diagram represents the document schema entities in the Visual DSA platform and the "
         "relationships between them. It outlines entity attributes, primary keys, foreign key references, and cardinality.",
         "documentation/diagrams/page5_er.png"),

        ("USE CASE DIAGRAM",
         "A Use Case Diagram shows how different actors interact with the system and represents the major functional operations "
         "available to each actor within the Visual DSA interactive platform.",
         "documentation/diagrams/page6_use_case.png"),

        ("CLASS DIAGRAM",
         "A Class Diagram represents the structure of the software system by showing components, models, controllers, and services. "
         "It illustrates how React client components, Express gateway controllers, and Mongoose document services collaborate.",
         "documentation/diagrams/page7_class.png"),

        ("SEQUENCE DIAGRAM",
         "A Sequence Diagram shows the sequence of interactions between system components over time. It explains how visualizer "
         "requests are initiated by the learner, verified by the React state machine, synchronized with pseudocode, dispatched "
         "to backend controllers, and committed to MongoDB.",
         "documentation/diagrams/page8_sequence.png")
    ]

    for title, desc, img_path in diagram_pages:
        add_heading(title, space_before=4, space_after=4)
        p_d = doc.add_paragraph(desc)
        p_d.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_d.paragraph_format.space_after = Pt(8)

        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_img = p_img.add_run()
            run_img.add_picture(img_path, width=Inches(5.7))
        
        doc.add_page_break()

    # =========================================================================
    # PAGES 9 to 13: APPLICATION SCREENSHOT PLATES (5 Distinct Pages, 2 per page)
    # =========================================================================
    screenshot_plates = [
        # Page 9: Plate I
        ("APPLICATION SCREENSHOTS — PLATE I", [
            ("The Platform Home and Curriculum Overview interface presents structured learning paths, topic cards, and difficulty badges.",
             "documentation/screenshots/01_home_dashboard.png"),
            ("The Array Visualizer interface renders dynamic memory element blocks, execution controls, and synchronized pseudocode.",
             "documentation/screenshots/02_array_visualizer.png")
        ]),
        # Page 10: Plate II
        ("APPLICATION SCREENSHOTS — PLATE II", [
            ("The Stack Visualizer interface illustrates Last-In First-Out (LIFO) operations with dynamic Top pointer tracking and overflow guards.",
             "documentation/screenshots/06_stack_visualizer.png"),
            ("The Linear Queue Visualizer demonstrates First-In First-Out (FIFO) pipeline execution with Front and Rear pointer management.",
             "documentation/screenshots/07_queue_visualizer.png")
        ]),
        # Page 11: Plate III
        ("APPLICATION SCREENSHOTS — PLATE III", [
            ("The Circular Queue Visualizer illustrates ring-buffer modulo wrap-around arithmetic ((rear + 1) % capacity) on a circular track.",
             "documentation/screenshots/08_circular_queue.png"),
            ("The Singly Linked List Visualizer renders non-contiguous memory blocks, dynamic next pointer reconnections, and traversal markers.",
             "documentation/screenshots/09_linkedlist_visualizer.png")
        ]),
        # Page 12: Plate IV
        ("APPLICATION SCREENSHOTS — PLATE IV", [
            ("The Sorting Playground visualizes comparison-based sorting algorithms (Bubble, Selection, Insertion, Merge, Quick) with color-coded comparison bars.",
             "documentation/screenshots/03_sorting_playground.png"),
            ("The Formative Practice Quiz interface offers topic-categorized multiple-choice questions with immediate explanation reveals.",
             "documentation/screenshots/04_practice_quiz.png")
        ]),
        # Page 13: Plate V
        ("APPLICATION SCREENSHOTS — PLATE V", [
            ("The Student Analytics Dashboard displays overall completion percentages, time spent, accuracy breakdowns, and topic recommendations.",
             "documentation/screenshots/10_progress_dashboard.png"),
            ("The Official Certification interface computes dynamic mastery thresholds and generates high-resolution HTML5 Canvas verified certificates.",
             "documentation/screenshots/05_certificate_page.png")
        ])
    ]

    for title, shots in screenshot_plates:
        add_heading(title, space_before=2, space_after=2)
        for desc, img_path in shots:
            p_desc = doc.add_paragraph(desc)
            p_desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_desc.paragraph_format.space_after = Pt(3)
            p_desc.paragraph_format.line_spacing = 1.05

            if os.path.exists(img_path):
                p_im = doc.add_paragraph()
                p_im.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_im.paragraph_format.space_after = Pt(6)
                r_im = p_im.add_run()
                r_im.add_picture(img_path, width=Inches(5.4))
        
        doc.add_page_break()

    # =========================================================================
    # PAGE 14: TEST REPORT & VERIFICATION ANALYSIS
    # =========================================================================
    add_heading("TEST REPORT", space_before=4, space_after=3)
    p_tr = doc.add_paragraph(
        "The Visual DSA interactive platform was rigorously tested to verify algorithmic correctness, data integrity, "
        "memory boundary safety rules, responsive UI transitions, and distributed API reliability. The comprehensive "
        "test cases executed across modules are documented below:"
    )
    p_tr.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_tr.paragraph_format.space_after = Pt(6)

    # Expanded 12 Test Cases Table
    table_test = doc.add_table(rows=13, cols=6)
    table_test.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_test, color="D3D3D3", sz="4")

    t_cols = ("ID", "Feature", "Test Case", "Expected Result", "Actual Result", "Status")
    for j, th in enumerate(t_cols):
        cell = table_test.rows[0].cells[j]
        cell.paragraphs[0].add_run(th).bold = True
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=50, bottom=50, left=50, right=50)

    test_data_12 = [
        ("TC01", "Array Ops", "Insert at index 2", "Elements shift right; value placed", "Shifted and placed accurately", "Pass"),
        ("TC02", "Stack Bounds", "Push on full stack", "Stack overflow warning prompt", "Overflow blocked with alert", "Pass"),
        ("TC03", "Queue Flow", "Dequeue vacant queue", "Queue underflow condition caught", "Underflow handled cleanly", "Pass"),
        ("TC04", "Circular Queue", "Enqueue past capacity-1", "Wrap-around via (rear+1)%cap", "Wrapped to index 0 accurately", "Pass"),
        ("TC05", "Linked List", "Insert at beginning", "Head pointer updated to new node", "Pointer updated in O(1) time", "Pass"),
        ("TC06", "Binary Search", "Target in sorted array", "Low/High converge to mid in O(log n)", "Found & highlighted in 3 steps", "Pass"),
        ("TC07", "Sorting Engine", "Quick sort unsorted array", "Lomuto partition organizes pivot", "Array sorted in ascending order", "Pass"),
        ("TC08", "Quiz Evaluator", "Submit valid quiz answer", "Score updated & explanation shown", "Explanation displayed with point", "Pass"),
        ("TC09", "Auth Guard", "Login invalid password", "HTTP 401 & failure counter incremented", "Rejected & counter incremented", "Pass"),
        ("TC10", "Progress API", "Upsert topic progress", "Atomic upsert in MongoDB collection", "201 Created & stored in DB", "Pass"),
        ("TC11", "AI Assistant", "Query Lemmy AI tutor", "Contextual response or local fallback", "Responded with Big-O analogy", "Pass"),
        ("TC12", "Canvas Cert", "Generate PNG certificate", "Canvas renders 1200x810 verified PNG", "Generated & downloaded clean PNG", "Pass")
    ]

    col_w = [Inches(0.6), Inches(1.1), Inches(1.4), Inches(1.6), Inches(1.4), Inches(0.6)]
    for idx, row_data in enumerate(test_data_12, start=1):
        row = table_test.rows[idx]
        for c_idx, val in enumerate(row_data):
            c = row.cells[c_idx]
            c.width = col_w[c_idx]
            run = c.paragraphs[0].add_run(val)
            run.font.size = Pt(8.5)
            if c_idx == 5:
                run.bold = True
                run.font.color.rgb = RGBColor(0x16, 0x65, 0x34) # Green
            set_cell_margins(c, top=40, bottom=40, left=50, right=50)

    p_tt_h = doc.add_paragraph()
    p_tt_h.paragraph_format.space_before = Pt(8)
    p_tt_h.paragraph_format.space_after = Pt(3)
    p_tt_h.add_run("Testing Techniques & Quality Assurance").bold = True

    p_tt = doc.add_paragraph(
        "Unit Testing: Component algorithmic functions including binary search partitions, Lomuto pivot swaps, circular queue "
        "modulo increments, and form regexes were individually verified.\n"
        "White Box Testing: Internal API gateway routing, JWT signature verification middleware, timing-safe OTP validation, "
        "and Mongoose compound-indexed upsert mutations were thoroughly audited.\n"
        "Black Box Testing: End-to-end user journeys from registration to certificate export were verified across desktop and mobile viewports.\n"
        "Smoke Testing: Live HTTP health-check verification (/api/health) and automated Playwright headless browser navigation "
        "were executed against the production deployment."
    )
    p_tt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_tt.paragraph_format.line_spacing = 1.15
    p_tt.paragraph_format.space_after = Pt(4)

    doc.add_page_break()

    # =========================================================================
    # PAGE 15: CONCLUSION AND FURTHER ENHANCEMENTS
    # =========================================================================
    add_heading("CONCLUSION AND FURTHER ENHANCEMENTS", space_before=4, space_after=6)
    
    p_c1 = doc.add_paragraph(
        "Visual DSA — Interactive Data Structures and Algorithms Visualization Platform successfully provides an end-to-end, "
        "reliable, and pedagogical solution for computer science education. Developed using modern full-stack engineering "
        "practices incorporating React 19, TypeScript, Node.js, Express.js, MongoDB, and Python, the platform addresses the "
        "long-standing educational challenge of visualizing abstract algorithmic memory transformations."
    )
    p_c1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_c1.paragraph_format.space_after = Pt(6)

    p_c2 = doc.add_paragraph(
        "Through dedicated visual sandboxes for linear data structures, searching, and sorting algorithms, students gain an "
        "intuitive understanding of pointer mutations, array shifts, and divide-and-conquer mechanics. The synchronization of "
        "live execution with line-by-line pseudocode bridges the cognitive gap between theoretical syntax and runtime memory behavior. "
        "Furthermore, formative retrieval quizzes, automated diagnostic progress analytics, and official verifiable certification "
        "motivate sustained engagement and conceptual mastery."
    )
    p_c2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_c2.paragraph_format.space_after = Pt(8)

    p_fe_title = doc.add_paragraph()
    p_fe_title.paragraph_format.space_before = Pt(6)
    p_fe_title.paragraph_format.space_after = Pt(4)
    r_fe = p_fe_title.add_run("Further Enhancements & Future Roadmap")
    r_fe.bold = True
    r_fe.font.size = Pt(11)

    future_points = [
        "Expansion to Non-Linear Hierarchical Structures: Incorporating Binary Search Trees (BST), AVL Tree self-balancing rotations, Red-Black Trees, and Min/Max Heaps.",
        "Graph Algorithm Visualizations: Developing interactive canvas graphs supporting Breadth-First Search (BFS), Depth-First Search (DFS), and Dijkstra's Shortest Path algorithm.",
        "Browser-Based Code Playground: Integrating an interactive Monaco/Ace code editor allowing students to write custom JavaScript or Python implementations and visualize their code dynamically.",
        "WebSocket Collaborative Classrooms: Implementing real-time multi-user synchronization so instructors can lead synchronized visual walkthroughs for remote student cohorts.",
        "Enhanced Security & Session Hardening: Migrating client JWT storage from localStorage to HTTP-only, secure, SameSite=Strict cookies to eliminate cross-site scripting exposure.",
        "Automated CI/CD Quality Pipeline: Embedding automated Vitest component regression tests and Supertest API integration suites into the GitHub Actions continuous integration workflow."
    ]

    for pt in future_points:
        p_pt = doc.add_paragraph(pt, style='List Bullet')
        p_pt.paragraph_format.space_after = Pt(4)
        p_pt.paragraph_format.line_spacing = 1.15
        p_pt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.save("documentation/Visual_DSA_Project_Documentation.docx")
    print("Created 15-page documentation/Visual_DSA_Project_Documentation.docx successfully")

if __name__ == '__main__':
    create_15_page_document()
