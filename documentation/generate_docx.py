import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

def add_page_border(section):
    secPr = section._sectPr
    pgBorders = parse_xml(
        f'<w:pgBorders {nsdecls("w")} w:offsetFrom="page">'
        f'  <w:top w:val="single" w:sz="12" w:space="24" w:color="000000"/>'
        f'  <w:left w:val="single" w:sz="12" w:space="24" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="12" w:space="24" w:color="000000"/>'
        f'  <w:right w:val="single" w:sz="12" w:space="24" w:color="000000"/>'
        f'</w:pgBorders>'
    )
    secPr.append(pgBorders)

def create_word_document():
    doc = Document()
    
    # Page setup - A4 with 0.75 inch margins
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    add_page_border(section)

    # Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(3)

    # Helper for headings
    def add_heading(text, space_before=6, space_after=4):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(12)
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        return p

    # =========================================================================
    # PAGE 1: COVER PAGE
    # =========================================================================
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(25)
    p_logo.paragraph_format.space_after = Pt(15)
    
    cit_logo_path = 'documentation/extracted_p1_img0_Image18.png'
    if os.path.exists(cit_logo_path):
        run_logo = p_logo.add_run()
        run_logo.add_picture(cit_logo_path, width=Inches(1.5))

    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("COIMBATORE INSTITUTE OF TECHNOLOGY—641014\nMSC – SOFTWARE SYSTEMS.")
    r_inst.bold = True
    r_inst.font.size = Pt(14)
    p_inst.paragraph_format.space_after = Pt(25)

    p_proj = doc.add_paragraph()
    p_proj.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_proj = p_proj.add_run("PROJECT TITLE : VISUAL DSA — INTERACTIVE DATA STRUCTURES AND ALGORITHMS VISUALIZATION PLATFORM.")
    r_proj.bold = True
    r_proj.font.size = Pt(13)
    p_proj.paragraph_format.space_after = Pt(8)

    p_mod = doc.add_paragraph()
    p_mod.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_mod = p_mod.add_run("MODULE : Interactive Memory Visualizer, Pseudocode Synchronization & AI Learning Assistant.")
    r_mod.font.size = Pt(12)
    p_mod.paragraph_format.space_after = Pt(45)

    # Student Details Table
    table_p1 = doc.add_table(rows=6, cols=2)
    table_p1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_p1, color="000000", sz="6")

    meta_rows = [
        ("NAME", "DHARANEESH V"),
        ("REGISTER NUMBER", "2503717672621011"),
        ("COURSE", "MSC – SOFTWARE SYSTEMS"),
        ("SUBJECT CODE", "25SS322"),
        ("SUBJECT", "PYTHON PROGRAMMING LABORATORY"),
        ("SEMESTER", "3")
    ]

    col_widths = [Inches(2.5), Inches(3.8)]
    for i, (k, v) in enumerate(meta_rows):
        row = table_p1.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = col_widths[0]
        c1.width = col_widths[1]
        set_cell_margins(c0, top=100, bottom=100, left=150, right=150)
        set_cell_margins(c1, top=100, bottom=100, left=150, right=150)
        
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(k)
        r0.bold = True
        r0.font.size = Pt(10.5)

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(v)
        r1.bold = True
        r1.font.size = Pt(10.5)

    doc.add_page_break()

    # =========================================================================
    # PAGE 2: SYNOPSIS, OBJECTIVE, AND SCOPE
    # =========================================================================
    add_heading("SYNOPSIS", space_before=4, space_after=3)
    p_syn1 = doc.add_paragraph(
        "Visual DSA is an interactive educational web-based platform developed to simplify the process of learning "
        "fundamental Data Structures and Algorithms. The platform replaces passive, static textbook code with responsive, "
        "step-by-step visual memory simulations, allowing learners to inspect pointer mutations, element displacements, "
        "and recursion mechanics in real time."
    )
    p_syn1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p_syn2 = doc.add_paragraph(
        "Module: Interactive Memory Visualizer, Pseudocode Synchronization & AI Learning Assistant is responsible for "
        "the core experiential learning lifecycle. It delivers dedicated visual sandboxes for linear arrays, singly linked "
        "lists, stacks (LIFO), queues (FIFO), circular queues (ring buffers), linear/binary searching, and five distinct "
        "sorting algorithms. Synchronized line-by-line pseudocode execution dynamically coordinates with graphic canvas state changes."
    )
    p_syn2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p_syn3 = doc.add_paragraph(
        "The application uses React 19 with Vite and TypeScript for the frontend client interface, and Node.js with "
        "Express.js for the primary API gateway. MongoDB is utilized for persistent cloud storage of student learning "
        "progress, operational activities, and quiz practice metrics. A supporting Python FastAPI microservice computes "
        "statistical mastery analytics and delivers SMTP password-reset OTP emails. An integrated Ollama-powered mascot tutor "
        "('Lemmy') provides context-aware tutoring and audio/speech synthesis."
    )
    p_syn3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    add_heading("OBJECTIVE", space_before=6, space_after=3)
    objectives = [
        "Render interactive, step-by-step memory visualizers for linear data structures, searching, and sorting.",
        "Synchronize algorithmic state transformations with highlighted, line-by-line pseudocode execution.",
        "Demonstrate pointer mutations, head/tail allocations, and modulo ring-buffer wrap-arounds graphically.",
        "Provide formative retrieval practice quizzes with immediate explanation reveals and accuracy metrics.",
        "Enforce defensive REST API security using JWT session tokens, bcrypt password hashing, and rate limiters.",
        "Deliver automated analytical diagnostic reports identifying student strengths, weaknesses, and study recommendations.",
        "Integrate a conversational local LLM AI tutor ('Lemmy') providing Big-O breakdowns and real-world analogies.",
        "Generate verifiable, high-resolution HTML5 Canvas Certificates of Completion with PNG and PDF exports."
    ]
    for obj in objectives:
        p_obj = doc.add_paragraph(obj, style='List Bullet')
        p_obj.paragraph_format.space_after = Pt(1.5)
        p_obj.paragraph_format.line_spacing = 1.1

    add_heading("SCOPE", space_before=6, space_after=3)
    scopes = [
        "Interactive visual sandboxes covering Arrays, Singly Linked Lists, Stacks, Queues, Circular Queues, Searching, and Sorting.",
        "Synchronous visualizer speed-control loop supporting custom operation parameters and boundary condition alerts.",
        "Dual-layer state persistence supporting authenticated MongoDB storage and seamless unauthenticated guest localStorage.",
        "FastAPI microservice executing statistical aggregation, mastery level classification, and SMTP email delivery.",
        "REST API gateway exposing validated, rate-limited endpoints for authentication, progress tracking, and quiz submissions.",
        "Client-side HTML5 Canvas rendering engine producing customized, verifiable completion certificates with unique serial IDs."
    ]
    for sc in scopes:
        p_sc = doc.add_paragraph(sc, style='List Bullet')
        p_sc.paragraph_format.space_after = Pt(1.5)
        p_sc.paragraph_format.line_spacing = 1.1

    doc.add_page_break()

    # =========================================================================
    # PAGE 3: PROJECT DESCRIPTION AND DB DESIGN (Calibrated to fit Page 3 exactly)
    # =========================================================================
    add_heading("PROJECT DESCRIPTION", space_before=2, space_after=2)
    p_pd1 = doc.add_paragraph(
        "Visual DSA is a full-stack, distributed web application consisting of a React 19 (Vite) frontend, a Node.js/Express.js "
        "API gateway, a MongoDB database, and a supporting Python FastAPI microservice. The core visualizer module delivers "
        "the complete interactive algorithmic learning lifecycle, connecting interactive simulations with verified cloud mastery records."
    )
    p_pd1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_pd1.paragraph_format.space_after = Pt(2)

    p_pd2 = doc.add_paragraph(
        "The frontend provides an intuitive multi-tab visualizer with playback speed sliders, synchronized pseudocode blocks, "
        "and conversational AI assistance. The Node.js backend manages secure REST API routing, input sanitization, rate limiting, "
        "and Mongoose persistence, while communicating with the Python microservice for heavy statistical analytics."
    )
    p_pd2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_pd2.paragraph_format.space_after = Pt(4)

    # Technology Table (6 rows)
    table_tech = doc.add_table(rows=7, cols=2)
    table_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_tech, color="D3D3D3", sz="4")

    h0, h1 = table_tech.rows[0].cells[0], table_tech.rows[0].cells[1]
    h0.paragraphs[0].add_run("Technology").bold = True
    h1.paragraphs[0].add_run("Purpose").bold = True
    set_cell_background(h0, "F1F5F9")
    set_cell_background(h1, "F1F5F9")
    set_cell_margins(h0, top=50, bottom=50, left=100, right=100)
    set_cell_margins(h1, top=50, bottom=50, left=100, right=100)

    tech_data = [
        ("React 19 & TypeScript", "Frontend user interface, component hierarchy, typed state models, and reactive DOM rendering"),
        ("Vite (v8)", "Fast Hot Module Replacement (HMR) build tooling and optimized production bundling"),
        ("Node.js & Express (v4)", "Asynchronous runtime executing REST API gateway, rate limiting, and security middleware"),
        ("MongoDB & Mongoose", "NoSQL document database and Object Data Modeling (ODM) defining schemas and validation"),
        ("Python 3 & FastAPI", "Supporting microservice for statistical mastery analytics, study time formatting, and SMTP OTP"),
        ("Ollama / Local LLM", "Contextual algorithmic tutor ('Lemmy') powered by llama3.2:3b with Web Audio and Speech synthesis")
    ]

    for idx, (t, p) in enumerate(tech_data, start=1):
        row = table_tech.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.5)
        run_t = c0.paragraphs[0].add_run(t)
        run_t.bold = True
        run_t.font.size = Pt(9.5)
        run_p = c1.paragraphs[0].add_run(p)
        run_p.font.size = Pt(9.5)
        set_cell_margins(c0, top=40, bottom=40, left=90, right=90)
        set_cell_margins(c1, top=40, bottom=40, left=90, right=90)

    add_heading("DB DESIGN", space_before=5, space_after=2)
    p_dbd = doc.add_paragraph(
        "Database Design describes how data is organized and stored in MongoDB. The system primarily interacts with the "
        "'users', 'progresses', 'practices', and 'activities' collections. The primary Progress schema structure is detailed below:"
    )
    p_dbd.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_dbd.paragraph_format.space_after = Pt(3)

    # DB Schema Table (7 data rows)
    table_db = doc.add_table(rows=8, cols=3)
    table_db.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_db, color="D3D3D3", sz="4")

    db_headers = ("Field", "Data Type", "Description")
    for j, dh in enumerate(db_headers):
        cell = table_db.rows[0].cells[j]
        cell.paragraphs[0].add_run(dh).bold = True
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=50, bottom=50, left=80, right=80)

    schema_data = [
        ("_id", "ObjectId", "Unique progress document identifier generated by MongoDB"),
        ("userId", "ObjectId", "Reference foreign key linking to the authenticated User document"),
        ("topic", "String", "Curriculum domain category (e.g., 'Array', 'Stack', 'CircularQueue')"),
        ("concept", "String", "Specific module concept name (e.g., 'Operations', 'Traversal')"),
        ("completed", "Boolean", "Completion flag indicating full mastery of the topic operations"),
        ("completionPercentage", "Number", "Calculated mastery progress metric ranging from 0 to 100"),
        ("timeSpent", "Number", "Cumulative study and visualizer engagement time in seconds")
    ]

    for idx, (f, dt, desc) in enumerate(schema_data, start=1):
        row = table_db.rows[idx]
        c0, c1, c2 = row.cells[0], row.cells[1], row.cells[2]
        c0.width = Inches(1.8)
        c1.width = Inches(1.2)
        c2.width = Inches(3.7)
        r0 = c0.paragraphs[0].add_run(f)
        r0.bold = True
        r0.font.size = Pt(9)
        r1 = c1.paragraphs[0].add_run(dt)
        r1.font.size = Pt(9)
        r2 = c2.paragraphs[0].add_run(desc)
        r2.font.size = Pt(9)
        set_cell_margins(c0, top=35, bottom=35, left=70, right=70)
        set_cell_margins(c1, top=35, bottom=35, left=70, right=70)
        set_cell_margins(c2, top=35, bottom=35, left=70, right=70)

    p_dbflow = doc.add_paragraph()
    p_dbflow.paragraph_format.space_before = Pt(4)
    p_dbflow.paragraph_format.space_after = Pt(2)
    r_flow = p_dbflow.add_run("Database flow: React Frontend -> Progress API (POST /api/progress) -> Node.js/Express -> Mongoose -> MongoDB -> Progresses & Activities Collections.")
    r_flow.bold = True
    r_flow.font.size = Pt(8.5)

    p_atom = doc.add_paragraph(
        "Data persistence guarantees consistency: upon visualizer operation completion, student records are atomically "
        "upserted using findOneAndUpdate() with compound index { userId: 1, topic: 1, concept: 1 }, cumulative time tallies are "
        "incremented, and verified activity audit events are logged in the Activity collection."
    )
    p_atom.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_atom.paragraph_format.line_spacing = 1.1
    p_atom.paragraph_format.space_after = Pt(2)

    doc.add_page_break()

    # =========================================================================
    # PAGES 4 to 8: DIAGRAMS (Each exactly 1 page)
    # =========================================================================
    diagram_pages = [
        ("DATA FLOW DIAGRAM (DFD)",
         "A Data Flow Diagram (DFD) represents how data moves through the Visual DSA interactive platform. It illustrates "
         "the interaction between learners, frontend visualizers, validation layers, Express API endpoints, MongoDB collections, "
         "the Python analysis service, and the local AI assistant runtime.",
         "documentation/diagrams/page4_dfd.png"),

        ("ER DIAGRAM",
         "An Entity Relationship (ER) Diagram represents the document schema entities in the Visual DSA platform and the "
         "relationships between them. It outlines entity attributes, primary keys, foreign key references, and cardinality.",
         "documentation/diagrams/page5_er.png"),

        ("USE CASE DIAGRAM",
         "A Use Case Diagram shows how different actors interact with the system and represents the major functional operations "
         "available to each actor within the Visual DSA interactive platform.",
         "documentation/diagrams/page6_use_case.png"),

        ("CLASS DIAGRAM",
         "A Class Diagram represents the structure of the software system by showing components, models, controllers, and services. "
         "It illustrates how React client components, Express gateway controllers, and Mongoose document services collaborate.",
         "documentation/diagrams/page7_class.png"),

        ("SEQUENCE DIAGRAM",
         "A Sequence Diagram shows the sequence of interactions between system components over time. It explains how visualizer "
         "requests are initiated by the learner, verified by the React state machine, synchronized with pseudocode, dispatched "
         "to backend controllers, and committed to MongoDB.",
         "documentation/diagrams/page8_sequence.png")
    ]

    for title, desc, img_path in diagram_pages:
        add_heading(title, space_before=4, space_after=4)
        p_d = doc.add_paragraph(desc)
        p_d.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_d.paragraph_format.space_after = Pt(8)

        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_img = p_img.add_run()
            run_img.add_picture(img_path, width=Inches(5.7))
        
        doc.add_page_break()

    # =========================================================================
    # PAGE 9: APPLICATION SCREENSHOTS (Exactly 1 page)
    # =========================================================================
    add_heading("APPLICATION SCREENSHOTS", space_before=2, space_after=3)
    
    # Screenshot 1: Array Visualizer
    p_sc1_desc = doc.add_paragraph(
        "The Array Visualizer interface renders dynamic memory element blocks, execution playback controls, bounds-checking alerts, "
        "and synchronized line-by-line pseudocode execution for insertions, deletions, traversals, and updates."
    )
    p_sc1_desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_sc1_desc.paragraph_format.space_after = Pt(4)

    img1_path = 'documentation/screenshots/02_array_visualizer.png'
    if os.path.exists(img1_path):
        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img1.paragraph_format.space_after = Pt(8)
        run_img1 = p_img1.add_run()
        run_img1.add_picture(img1_path, width=Inches(5.4))

    # Screenshot 2: Certificate Page
    p_sc2_desc = doc.add_paragraph(
        "The Visual DSA Certification interface computes dynamic topic completion metrics, offers custom recipient name editing, "
        "and provides verified HTML5 Canvas rendering for instantaneous high-resolution PNG downloads and print-to-PDF generation."
    )
    p_sc2_desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_sc2_desc.paragraph_format.space_after = Pt(4)

    img2_path = 'documentation/screenshots/05_certificate_page.png'
    if os.path.exists(img2_path):
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_img2 = p_img2.add_run()
        run_img2.add_picture(img2_path, width=Inches(5.4))

    doc.add_page_break()

    # =========================================================================
    # PAGE 10: TEST REPORT, CONCLUSION AND FURTHER ENHANCEMENTS (Calibrated)
    # =========================================================================
    add_heading("TEST REPORT", space_before=2, space_after=2)
    p_tr = doc.add_paragraph(
        "The Visual DSA interactive platform was thoroughly tested to verify algorithmic correctness, data integrity, "
        "memory boundary safety rules, responsive state transitions, and database persistence reliability."
    )
    p_tr.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_tr.paragraph_format.space_after = Pt(3)

    # Test Table (10 test cases)
    table_test = doc.add_table(rows=11, cols=6)
    table_test.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_test, color="D3D3D3", sz="4")

    t_cols = ("ID", "Feature", "Test Case", "Expected", "Actual", "Status")
    for j, th in enumerate(t_cols):
        cell = table_test.rows[0].cells[j]
        cell.paragraphs[0].add_run(th).bold = True
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=40, bottom=40, left=50, right=50)

    test_data = [
        ("TC01", "Array Ops", "Insert at index 2", "Elements shift right; new item placed", "Shifted and placed accurately", "Pass"),
        ("TC02", "Stack Bounds", "Push on full stack", "Stack overflow warning prompt", "Overflow blocked with alert", "Pass"),
        ("TC03", "Queue Flow", "Dequeue vacant queue", "Queue underflow condition caught", "Underflow handled cleanly", "Pass"),
        ("TC04", "Circular Queue", "Enqueue past capacity-1", "Index wraps around via (rear+1)%cap", "Wrapped to index 0 accurately", "Pass"),
        ("TC05", "Binary Search", "Target item in sorted array", "Low/High converge to mid in O(log n)", "Found & highlighted in 3 steps", "Pass"),
        ("TC06", "Sorting Engine", "Sort unsorted array (Quick)", "Lomuto partition organizes around pivot", "Array sorted in ascending order", "Pass"),
        ("TC07", "Quiz Evaluator", "Submit valid quiz answer", "Instant score update & explanation reveal", "Explanation displayed with point", "Pass"),
        ("TC08", "Auth Guard", "Login invalid password", "HTTP 401 & failure count incremented", "Rejected & counter incremented", "Pass"),
        ("TC09", "Progress API", "Upsert topic progress", "Atomic upsert in MongoDB collection", "201 Created & stored in DB", "Pass"),
        ("TC10", "Canvas Cert", "Generate PNG certificate", "Canvas renders 1200x810 verified PNG", "Generated & downloaded clean PNG", "Pass")
    ]

    col_w = [Inches(0.6), Inches(1.1), Inches(1.4), Inches(1.6), Inches(1.4), Inches(0.6)]
    for idx, row_data in enumerate(test_data, start=1):
        row = table_test.rows[idx]
        for c_idx, val in enumerate(row_data):
            c = row.cells[c_idx]
            c.width = col_w[c_idx]
            run = c.paragraphs[0].add_run(val)
            run.font.size = Pt(8)
            if c_idx == 5:
                run.bold = True
                run.font.color.rgb = RGBColor(0x16, 0x65, 0x34) # Green
            set_cell_margins(c, top=30, bottom=30, left=45, right=45)

    # Testing Techniques (Compact)
    p_tt_h = doc.add_paragraph()
    p_tt_h.paragraph_format.space_before = Pt(4)
    p_tt_h.paragraph_format.space_after = Pt(1)
    p_tt_h.add_run("Testing Techniques").bold = True

    p_tt = doc.add_paragraph(
        "Unit Testing: Component algorithmic functions including binary search partitions, Lomuto pivot swaps, circular queue modulo increments, and form regexes were individually verified. "
        "White Box Testing: Internal API gateway routing, JWT signature verification middleware, timing-safe OTP validation, and Mongoose compound-indexed upsert mutations were examined. "
        "Black Box Testing: End-to-end user journeys from registration to certificate export were verified across desktop and mobile viewports. "
        "Smoke Testing: Live HTTP health-check verification (/api/health) and automated Playwright headless browser navigation were executed on production endpoints."
    )
    p_tt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_tt.paragraph_format.line_spacing = 1.05
    p_tt.paragraph_format.space_after = Pt(3)

    # Conclusion & Further Enhancements (Compact)
    add_heading("CONCLUSION AND FURTHER ENHANCEMENTS", space_before=4, space_after=2)
    p_conc = doc.add_paragraph(
        "Visual DSA successfully provides an end-to-end, reliable solution for interactive computer science education. "
        "Developed using React 19, TypeScript, Node.js, Express.js, MongoDB, and Python, the platform features real-time "
        "graphical memory animations, synchronized pseudocode execution, formative quiz evaluations, and verifiable certificate issuance."
    )
    p_conc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_conc.paragraph_format.line_spacing = 1.05
    p_conc.paragraph_format.space_after = Pt(2)

    p_enh_h = doc.add_paragraph()
    p_enh_h.paragraph_format.space_before = Pt(2)
    p_enh_h.paragraph_format.space_after = Pt(1)
    p_enh_h.add_run("Further Enhancements").bold = True

    enhancements = [
        "Expansion to Non-Linear Data Structures including Binary Search Trees (BST), AVL Trees, and Graph traversals (BFS/DFS).",
        "Implementation of a browser-based live code editor sandbox allowing students to visualize their custom algorithm code.",
        "Transition of JWT token storage from localStorage to HTTP-only, secure, SameSite=Strict cookies for enhanced security.",
        "WebSocket-enabled collaborative rooms allowing multiple students and instructors to step through visualizers simultaneously.",
        "Automated continuous integration (CI) test suite incorporating Vitest component tests and Supertest API regression suites."
    ]
    for enh in enhancements:
        p_enh = doc.add_paragraph(enh, style='List Bullet')
        p_enh.paragraph_format.space_after = Pt(1)
        p_enh.paragraph_format.line_spacing = 1.05

    doc.save("documentation/Visual_DSA_Project_Documentation.docx")
    print("Recalibrated documentation/Visual_DSA_Project_Documentation.docx saved successfully")

if __name__ == '__main__':
    create_word_document()
